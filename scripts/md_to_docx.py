#!/usr/bin/env python3
"""Convert the Executive Platform markdown packages to Word for the funding file."""

from __future__ import annotations

import re
from pathlib import Path

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Inches, Pt, RGBColor, Cm, Emu
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image

IMG_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")

NAVY = RGBColor(0x10, 0x20, 0x33)
RED = RGBColor(0x9F, 0x1B, 0x2C)
SLATE = RGBColor(0x33, 0x3A, 0x44)
GREY = RGBColor(0x5B, 0x64, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = "F4F6F8"
HEADER_BG = "102033"

ROOT = Path("/Users/chilr357/L-Bax_Archive/Thousand Oaks DDT/2026 Strategy/Dean Capability Tracker")

FILES = [
    (
        ROOT / "GSQ-DDT-Executive-Platform-Design-Blueprint.md",
        ROOT / "GSQ-DDT-Executive-Platform-Design-Blueprint.docx",
        "Business Design Package",
    ),
    (
        ROOT / "GSQ-DDT-Executive-Platform-Executive-Summary.md",
        ROOT / "GSQ-DDT-Executive-Platform-Executive-Summary.docx",
        "Executive Summary",
    ),
    (
        ROOT / "GSQ-DDT-Executive-Platform-TakOS-Onboarding-Readiness.md",
        ROOT / "GSQ-DDT-Executive-Platform-TakOS-Onboarding-Readiness.docx",
        "TakOS Onboarding Readiness",
    ),
    (
        ROOT / "GSQ-DDT-Executive-Platform-Executive-Design.md",
        ROOT / "GSQ-DDT-Executive-Platform-Executive-Design.docx",
        "Executive Design Extract",
    ),
    (
        ROOT / "GSQ-DDT-Executive-Platform-TakOS-Submission-Checklist.md",
        ROOT / "GSQ-DDT-Executive-Platform-TakOS-Submission-Checklist.docx",
        "TakOS Submission Checklist",
    ),
    (
        ROOT / "GSQ-DDT-Executive-Platform-SOW-Altimetrik-TM-Oct-Dec-2026.md",
        ROOT / "GSQ-DDT-Executive-Platform-SOW-Altimetrik-TM-Oct-Dec-2026.docx",
        "Altimetrik T&M SOW draft",
    ),
    (
        ROOT / "GSQ-DDT-Executive-Platform-PowerApps-SharePoint-Prototype.md",
        ROOT / "GSQ-DDT-Executive-Platform-PowerApps-SharePoint-Prototype.docx",
        "Power Apps + SharePoint prototype",
    ),
]


def set_run_font(run, name="Calibri", size=11, color=SLATE, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D0D5DD")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_text(cell, text: str, *, header=False, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.strip())
    if header:
        set_run_font(run, size=9, color=WHITE, bold=True)
    else:
        set_run_font(run, size=9, color=SLATE, bold=bold)
    set_cell_borders(cell)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = SLATE
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, space_before, space_after in [
        ("Heading 1", 18, NAVY, 18, 8),
        ("Heading 2", 14, RED, 16, 6),
        ("Heading 3", 12, NAVY, 12, 4),
        ("Heading 4", 11, NAVY, 10, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(space_before)
        st.paragraph_format.space_after = Pt(space_after)

    if "Code Block" not in [s.name for s in styles]:
        cs = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        cs.font.name = "Consolas"
        cs.font.size = Pt(8)
        cs.font.color.rgb = NAVY
        cs.paragraph_format.space_before = Pt(4)
        cs.paragraph_format.space_after = Pt(4)
        cs.paragraph_format.left_indent = Cm(0.4)


def add_page_header(doc: Document, subtitle: str):
    header = doc.sections[0].header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.text = ""
    r = p.add_run("GSQ GML DD&T  ·  Executive Platform  ·  ")
    set_run_font(r, size=8, color=GREY)
    r2 = p.add_run(subtitle)
    set_run_font(r2, size=8, color=RED, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fr = fp.add_run("Confidential — Takeda internal  ·  September 2026  ·  Page ")
    set_run_font(fr, size=8, color=GREY)
    # page number field
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run = fp.add_run()
    run._r.append(fld)
    run2 = fp.add_run()
    run2._r.append(instr)
    run3 = fp.add_run()
    run3._r.append(fld2)
    set_run_font(run, size=8, color=GREY)
    set_run_font(run2, size=8, color=GREY)
    set_run_font(run3, size=8, color=GREY)


def add_inline_runs(paragraph, text: str, *, size=11, color=SLATE, italic=False):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, color=color, bold=True, italic=italic)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=size - 1, color=NAVY)
        elif part.startswith("[") and "](" in part:
            label = part[1 : part.index("]")]
            run = paragraph.add_run(label)
            set_run_font(run, size=size, color=RED)
            run.underline = True
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color, italic=italic)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        if re.match(r"^\s*\|?\s*-{2,}", line.replace("|", " | ")):
            # separator
            if set(line.replace("|", "").replace(":", "").replace("-", "").strip()) == set():
                continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def add_figure(doc: Document, caption: str, rel_path: str):
    src = (ROOT / rel_path).resolve()
    if not src.exists():
        p = doc.add_paragraph()
        add_inline_runs(p, f"[Missing figure: {rel_path}]", size=10, color=RED)
        return
    img = Image.open(src)
    if img.mode in ("RGBA", "P"):
        img = img.convert("RGB")
    max_w, max_h = 1600, 1200
    ratio = min(max_w / img.width, max_h / img.height, 1.0)
    if ratio < 1.0:
        img = img.resize((int(img.width * ratio), int(img.height * ratio)), Image.Resampling.LANCZOS)
    buf = BytesIO()
    buf.name = "figure.jpg"
    img.save(buf, format="JPEG", quality=82, optimize=True)
    buf.seek(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    # Keep tall screenshots (One pager) on one page; landscape shots stay full width.
    display_w = 6.5 if img.width >= img.height else 5.8
    run.add_picture(buf, width=Inches(display_w))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.space_after = Pt(12)
    add_inline_runs(cap, caption, size=9, color=GREY, italic=True)


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            val = row[j] if j < len(row) else ""
            val = re.sub(r"\*\*([^*]+)\*\*", r"\1", val)
            val = re.sub(r"`([^`]+)`", r"\1", val)
            val = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", val)
            if i == 0:
                shade_cell(cell, HEADER_BG)
                set_cell_text(cell, val, header=True)
            else:
                if i % 2 == 0:
                    shade_cell(cell, ROW_ALT)
                set_cell_text(cell, val, bold=j == 0)
    doc.add_paragraph()


def convert(md_path: Path, docx_path: Path, subtitle: str):
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    configure_styles(doc)
    add_page_header(doc, subtitle)

    i = 0
    in_code = False
    code_buf: list[str] = []
    title_done = False

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                p = doc.add_paragraph("\n".join(code_buf), style="Code Block")
                p.paragraph_format.left_indent = Cm(0.3)
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            run = p.add_run(line[2:].strip())
            set_run_font(run, size=22, color=NAVY, bold=True)
            p.paragraph_format.space_after = Pt(2)
            title_done = True
            i += 1
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
            i += 1
            continue

        if line.strip().startswith("|") and i + 1 < len(lines) and "|" in lines[i + 1]:
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if re.match(r"^\s*\|?\s*:?-{2,}", lines[i].replace(" ", "")):
                    i += 1
                    continue
                block.append(lines[i])
                i += 1
            add_table(doc, parse_table(block))
            continue

        if re.match(r"^\s*[-*] ", line):
            text_item = re.sub(r"^\s*[-*] ", "", line)
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, text_item, size=11)
            i += 1
            continue

        if re.match(r"^\s*\d+\. ", line):
            text_item = re.sub(r"^\s*\d+\. ", "", line)
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, text_item, size=11)
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        img = IMG_RE.match(line.strip())
        if img:
            add_figure(doc, img.group(1), img.group(2))
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, line.strip(), size=11)
        i += 1

    doc.save(docx_path)
    print(f"Wrote {docx_path.name} ({docx_path.stat().st_size:,} bytes)")


def main():
    for src, dest, subtitle in FILES:
        convert(src, dest, subtitle)


if __name__ == "__main__":
    main()
